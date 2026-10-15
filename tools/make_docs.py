"""把 docs/*.md 转成 Word 文档（便于直接提交）。

用法::

    python tools/make_docs.py                     # 转换 docs 下所有 .md
    python tools/make_docs.py --only 亮点说明      # 只转换指定文件（不含扩展名）

依赖 ``python-docx``（pip install python-docx）。支持：标题、段落、无序列表、
有序列表、代码块、Markdown 表格、粗体行内标记（**text**）。
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DOCS_DIR = PROJECT_ROOT / "docs"

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor
except ImportError:  # pragma: no cover
    print("需要 python-docx：pip install python-docx")
    sys.exit(1)


INLINE_BOLD = re.compile(r"\*\*(.+?)\*\*")


def add_inline(paragraph, text: str):
    position = 0
    for match in INLINE_BOLD.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position : match.start()])
        run = paragraph.add_run(match.group(1))
        run.bold = True
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def add_code_block(document: Document, lines: list[str]):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Pt(16)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run("\n".join(lines))
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x1F, 0x3B, 0x73)


def add_table(document: Document, rows: list[list[str]]):
    if not rows:
        return
    table = document.add_table(rows=0, cols=len(rows[0]))
    table.style = "Light Grid Accent 1"
    for index, row in enumerate(rows):
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = ""
            paragraph = cell.paragraphs[0]
            add_inline(paragraph, text.strip())
            for run in paragraph.runs:
                run.font.size = Pt(9)
                if index == 0:
                    run.bold = True


def convert(source: Path, target: Path):
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10.5)

    lines = source.read_text(encoding="utf-8").splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if stripped.startswith("```"):
            index += 1
            block: list[str] = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                block.append(lines[index])
                index += 1
            add_code_block(document, block)
        elif re.match(r"^#{1,4}\s", stripped):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped[level:].strip()
            heading = document.add_heading(level=min(level, 4))
            add_inline(heading, text)
        elif stripped.startswith("|"):
            rows: list[list[str]] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                raw = lines[index].strip().strip("|")
                if not re.match(r"^[\s:\-|]+$", raw):
                    rows.append(raw.split("|"))
                index += 1
            add_table(document, rows)
            continue
        elif re.match(r"^\s*[-*]\s+", line):
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, re.sub(r"^\s*[-*]\s+", "", line))
        elif re.match(r"^\s*\d+\.\s+", line):
            paragraph = document.add_paragraph(style="List Number")
            add_inline(paragraph, re.sub(r"^\s*\d+\.\s+", "", line))
        elif stripped.startswith(">"):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Pt(14)
            run_text = stripped.lstrip(">").strip()
            add_inline(paragraph, run_text)
            for run in paragraph.runs:
                run.italic = True
                run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        elif stripped == "---":
            document.add_paragraph()
        elif stripped:
            paragraph = document.add_paragraph()
            add_inline(paragraph, stripped)
        index += 1

    document.save(target)
    print(f"已生成 {target.relative_to(PROJECT_ROOT)}")


def main() -> int:
    parser = argparse.ArgumentParser(description="Markdown → Word 转换")
    parser.add_argument("--only", default=None, help="只转换指定文件名（不含扩展名）")
    args = parser.parse_args()

    sources = sorted(DOCS_DIR.glob("*.md"))
    if args.only:
        sources = [path for path in sources if path.stem == args.only]
    if not sources:
        print("没有找到要转换的 Markdown 文件")
        return 1

    succeeded: list[str] = []
    failed: list[tuple[str, str]] = []
    for source in sources:
        target = source.with_suffix(".docx")
        try:
            convert(source, target)
            succeeded.append(target.name)
        except PermissionError:
            failed.append((target.name, "文件正被 Word / WPS 打开，请关闭后重试"))
        except Exception as exc:  # noqa: BLE001
            failed.append((target.name, f"{type(exc).__name__}: {exc}"))

    print()
    print(f"成功 {len(succeeded)} 个：{', '.join(succeeded) if succeeded else '无'}")
    for name, reason in failed:
        print(f"失败：{name} —— {reason}")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
